"""Complete Excel, CSV, and Markdown exports from the audited SQLite ledger."""

from __future__ import annotations

import csv
import json
import sqlite3
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


EXCEL_CELL_LIMIT = 32767
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


@dataclass(slots=True)
class ExportStats:
    facts: int
    evidence_rows: int
    workbook_path: Path
    csv_path: Path
    markdown_path: Path


def _safe_cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value if len(value) <= EXCEL_CELL_LIMIT else value[:32748] + "… [truncated]"


def _configure_sheet(ws, widths: dict[str, float], table_name: str | None = None) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    if table_name and ws.max_row >= 2:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)


def _evidence_records(conn: sqlite3.Connection, ids: list[str]) -> list[sqlite3.Row]:
    if not ids:
        return []
    placeholders = ",".join("?" for _ in ids)
    rows = conn.execute(
        f"""
        SELECT rf.*, p.page_label FROM raw_facts rf
        JOIN chunks c ON c.id=rf.chunk_id JOIN pages p ON p.id=c.page_id
        WHERE rf.id IN ({placeholders})
        """,
        tuple(ids),
    ).fetchall()
    by_id = {row["id"]: row for row in rows}
    return [by_id[item] for item in ids if item in by_id]


def _evidence_text(conn: sqlite3.Connection, ids: list[str]) -> str:
    items = []
    for row in _evidence_records(conn, ids):
        label = row["page_label"] or str(row["page_number"])
        section = f", {row['chapter_section']}" if row["chapter_section"] else ""
        items.append(
            f"{row['evidence_class']} — PDF p. {row['page_number']} "
            f"(printed {label}{section}): “{row['evidence_quote']}”"
        )
    return "\n\n".join(items)


def _quote_block(text: str) -> str:
    return "\n".join(f"> {line}" for line in text.splitlines())


def _write_markdown(
    conn: sqlite3.Connection,
    path: Path,
    facts: list[sqlite3.Row],
    sources: dict[str, sqlite3.Row],
    project_title: str,
) -> None:
    lines = [
        f"# {project_title}: Fact Consensus Report",
        "",
        f"Generated: {datetime.now(timezone.utc).isoformat()}",
        "",
        "> The numerical evidence index is a reproducible research/triage measure, not a",
        "> calibrated probability that a claim is true. Quotations are verified against",
        "> extracted PDF page text; historical interpretation still requires human review.",
        "",
        "## Mathematical method",
        "",
        "For source *i*, source quality `q_i` is the policy-weighted geometric mean of",
        "authority, firsthand access, temporal proximity, documentary basis, independence,",
        "specialist expertise, and bias safety (`1 − political_bias_risk`). Evidence is",
        "multiplied by its E1–E7 class, extraction confidence, alignment confidence, and",
        "categorical stance value. Within each dependence cluster only the strongest",
        "support and strongest D1 contradiction count.",
        "",
        "`I(F) = 100 × (A − D) / (1 + A + D)`",
        "",
        "`E(F) = 100 × (1 − exp(−(A + D)))`",
        "",
        "Here `A` is independent support mass and `D` is penalizable D1 contradiction",
        "mass. D2–D9 disagreements remain visible but receive no automatic penalty.",
        "",
        "## Summary",
        "",
    ]
    counts: dict[str, int] = {}
    for fact in facts:
        counts[fact["grade"]] = counts.get(fact["grade"], 0) + 1
    lines.extend(f"- {label}: {count}" for label, count in sorted(counts.items()))
    lines.extend(["", f"- Total canonical claims: {len(facts)}", "", "---", ""])

    for position, fact in enumerate(facts, start=1):
        fact_label = f"F{position:06d}"
        lines.extend(
            [
                f"## {fact_label} — {fact['canonical_claim']}",
                "",
                f"- Database ID: `{fact['id']}`",
                f"- Event: {fact['event'] or 'Unclassified'}",
                f"- Taxonomy: `{fact['taxonomy_path'] or 'Unclassified'}`",
                f"- Claim type: `{fact['claim_type']}`",
                f"- Source date: {fact['date_source'] or 'Not stated'}",
                f"- Normalized date: {fact['date_normalized'] or 'Not established'}",
                f"- Location: {fact['location'] or 'Not stated'}",
                f"- Classification: **{fact['grade']}**",
                f"- Confidence tier: **{fact['confidence_tier']}**",
                f"- Evidence direction index: {fact['research_index']:.2f} (−100 to +100)",
                f"- Evidence strength: {fact['evidence_strength']:.2f}%",
                f"- Coverage of independent source quality: {100*fact['evidence_coverage']:.2f}%",
                f"- Independent/dependent support: {fact['independent_support_count']} / {fact['dependent_support_count']}",
                f"- Supporting E1/E2/E4 source counts: {fact['primary_document_count']} / {fact['eyewitness_count']} / {fact['scholarly_reconstruction_count']}",
                "",
            ]
        )
        decisions = conn.execute(
            "SELECT * FROM stances WHERE canonical_id=? ORDER BY source_id", (fact["id"],)
        ).fetchall()
        evidentiary = [row for row in decisions if row["stance"] != "silent"]
        lines.extend(["### Source evidence", ""])
        if not evidentiary:
            lines.extend(["No supporting or contradictory passage was verified.", ""])
        for decision in evidentiary:
            source = sources[decision["source_id"]]
            contradiction = (
                f"; {decision['contradiction_type']}" if decision["contradiction_type"] else ""
            )
            lines.extend(
                [
                    f"#### {source['title']} (`{source['id']}`)",
                    "",
                    f"Stance: **{decision['stance']}**; coverage: `{decision['coverage']}`{contradiction}.",
                    "",
                ]
            )
            for heading, json_ids in (
                ("Supporting passage", decision["support_evidence_ids_json"]),
                ("Contradicting/discrepant passage", decision["contradiction_evidence_ids_json"]),
            ):
                records = _evidence_records(conn, json.loads(json_ids))
                for record in records:
                    label = record["page_label"] or str(record["page_number"])
                    lines.extend(
                        [
                            f"**{heading} — {record['evidence_class']}, PDF p. {record['page_number']} (printed {label})**",
                            "",
                            _quote_block(record["evidence_quote"]),
                            "",
                        ]
                    )
            lines.extend([f"Rationale: {decision['rationale']}", ""])
            if decision["chronology_notes"]:
                lines.extend([f"Chronology note: {decision['chronology_notes']}", ""])
            if decision["scope_notes"]:
                lines.extend([f"Scope note: {decision['scope_notes']}", ""])

        covered_silent = [
            sources[row["source_id"]]["title"] for row in decisions
            if row["coverage"] == "covered_but_silent"
        ]
        not_covered = [
            sources[row["source_id"]]["title"] for row in decisions
            if row["coverage"] == "not_covered"
        ]
        lines.extend(
            [
                "### Silence and coverage",
                "",
                "- Covered but silent: " + (", ".join(covered_silent) or "None"),
                "- Not covered/retrieved: " + (", ".join(not_covered) or "None"),
                "",
                "---",
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def export_ledger(
    conn: sqlite3.Connection,
    output_dir: Path,
    *,
    project_title: str = "Historical Fact Consensus",
) -> ExportStats:
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = output_dir / "fact_consensus_matrix.xlsx"
    csv_path = output_dir / "fact_consensus_matrix.csv"
    markdown_path = output_dir / "fact_consensus_report.md"
    facts = conn.execute(
        """
        SELECT cf.*, sc.* FROM canonical_facts cf JOIN scores sc ON sc.canonical_id=cf.id
        ORDER BY CASE sc.grade WHEN 'FACT_CONTESTED' THEN 0 WHEN 'FACT_ESTABLISHED' THEN 1
                 WHEN 'FACT_PROBABLE' THEN 2 WHEN 'FACT_PLAUSIBLE' THEN 3 ELSE 4 END,
                 sc.evidence_strength DESC, cf.canonical_claim
        """
    ).fetchall()
    if not facts:
        raise RuntimeError("No scored facts found. Run score before export.")
    source_rows = conn.execute("SELECT * FROM sources ORDER BY id").fetchall()
    sources = {row["id"]: row for row in source_rows}

    headers = [
        "Fact ID", "Event", "Taxonomy", "Canonical atomic claim", "Claim type",
        "Source date", "Normalized date", "Calendar", "Location", "Classification",
        "Confidence tier", "Evidence direction index", "Evidence strength %",
        "Coverage %", "Conflict index %", "Independent support", "Dependent support",
        "Supporting E1 documents", "Supporting E2 eyewitnesses",
        "Supporting E4 scholarship", "Serious D1 contradictions",
        "Explicit support", "Implicit support", "Contradiction/discrepancy", "Mixed",
        "Covered but silent", "Not covered",
    ]
    matrix: list[list[object]] = []
    details: list[list[object]] = []
    for position, fact in enumerate(facts, start=1):
        decisions = conn.execute(
            "SELECT * FROM stances WHERE canonical_id=? ORDER BY source_id", (fact["id"],)
        ).fetchall()
        groups = {
            "explicit_support": [], "implicit_support": [], "explicit_contradiction": [],
            "mixed": [], "covered_but_silent": [], "not_covered": [],
        }
        for decision in decisions:
            source = sources[decision["source_id"]]
            if decision["stance"] == "silent":
                groups[decision["coverage"]].append(source["title"])
            else:
                groups[decision["stance"]].append(source["title"])
            support_ids = json.loads(decision["support_evidence_ids_json"])
            contradiction_ids = json.loads(decision["contradiction_evidence_ids_json"])
            details.append(
                [
                    f"F{position:06d}", fact["id"],
                    fact["canonical_claim"], source["id"], source["title"],
                    source["independence_cluster_id"], source["credibility_weight"],
                    decision["coverage"], decision["stance"], decision["contradiction_type"] or "",
                    decision["confidence"], _evidence_text(conn, support_ids),
                    _evidence_text(conn, contradiction_ids), decision["rationale"],
                    decision["chronology_notes"] or "", decision["scope_notes"] or "",
                ]
            )
        matrix.append(
            [
                f"F{position:06d}", fact["event"] or "", fact["taxonomy_path"] or "",
                fact["canonical_claim"], fact["claim_type"], fact["date_source"] or "",
                fact["date_normalized"] or "", fact["calendar"] or "", fact["location"] or "",
                fact["grade"], fact["confidence_tier"], fact["research_index"],
                fact["evidence_strength"] / 100, fact["evidence_coverage"], fact["conflict_index"],
                fact["independent_support_count"], fact["dependent_support_count"],
                fact["primary_document_count"], fact["eyewitness_count"],
                fact["scholarly_reconstruction_count"], fact["serious_contradiction_count"],
                "\n".join(groups["explicit_support"]), "\n".join(groups["implicit_support"]),
                "\n".join(groups["explicit_contradiction"]), "\n".join(groups["mixed"]),
                "\n".join(groups["covered_but_silent"]), "\n".join(groups["not_covered"]),
            ]
        )

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Fact Ledger"
    ws.append(headers)
    for row in matrix:
        ws.append([_safe_cell(value) for value in row])
    for row_number in range(2, ws.max_row + 1):
        for column in (13, 14, 15):
            ws.cell(row_number, column).number_format = "0.0%"
    ws.conditional_formatting.add(
        f"L2:L{ws.max_row}",
        ColorScaleRule(start_type="num", start_value=-100, start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFEB84",
                       end_type="num", end_value=100, end_color="63BE7B"),
    )
    _configure_sheet(ws, {"A": 12, "B": 24, "C": 30, "D": 70, "E": 22,
                          "F": 18, "G": 18, "H": 14, "I": 20, "J": 24,
                          "K": 38, "L": 20, "M": 18, "N": 14, "O": 16,
                          "P": 14, "Q": 14, "R": 12, "S": 14, "T": 14,
                          "U": 18, "V": 32, "W": 32, "X": 35, "Y": 30,
                          "Z": 32, "AA": 32}, "FactLedgerTable")

    detail_headers = [
        "Fact ID", "Database ID", "Canonical claim", "Source ID", "Source title",
        "Independence cluster", "Source quality q_i", "Coverage", "Stance",
        "Disagreement type", "Alignment confidence", "Supporting evidence",
        "Contradicting/discrepant evidence", "Rationale", "Chronology notes", "Scope notes",
    ]
    evidence_ws = workbook.create_sheet("Source Evidence")
    evidence_ws.append(detail_headers)
    for row in details:
        evidence_ws.append([_safe_cell(value) for value in row])
    _configure_sheet(evidence_ws, {"A": 12, "B": 34, "C": 65, "D": 30, "E": 40,
                                   "F": 30, "G": 16, "H": 20, "I": 23, "J": 22,
                                   "K": 18, "L": 80, "M": 80, "N": 55, "O": 35,
                                   "P": 35}, "SourceEvidenceTable")

    source_ws = workbook.create_sheet("Source Policy")
    source_headers = [
        "Source ID", "Filename", "Title", "Author", "Role", "Default evidence class",
        "Independence cluster", "Computed q_i", "Authority", "Firsthand", "Temporal",
        "Documentary", "Independence", "Expertise", "Political bias risk",
        "Description", "Bias notes", "PDF metadata", "SHA-256", "PDF pages",
    ]
    source_ws.append(source_headers)
    for source in source_rows:
        source_ws.append(
            [
                _safe_cell(value)
                for value in [
                    source["id"], source["filename"], source["title"],
                    source["author"] or "", source["role"],
                    source["evidence_class_default"], source["independence_cluster_id"],
                    source["credibility_weight"], source["source_authority"],
                    source["firsthand_access"], source["temporal_proximity"],
                    source["documentary_basis"], source["independence"],
                    source["specialist_expertise"], source["political_bias_risk"],
                    source["description"], source["bias_notes"],
                    source["pdf_metadata_json"] or "", source["file_sha256"] or "",
                    source["page_count"] or 0,
                ]
            ]
        )
    _configure_sheet(source_ws, {"A": 30, "B": 48, "C": 42, "D": 24, "E": 16,
                                 "F": 30, "G": 30, "H": 14, "I": 12, "J": 12,
                                 "K": 12, "L": 12, "M": 14, "N": 12, "O": 17,
                                 "P": 55, "Q": 55, "R": 55, "S": 66, "T": 12},
                     "SourcePolicyTable")

    method_ws = workbook.create_sheet("Methodology")
    method_ws.append(["Item", "Definition"])
    for item in [
        ("Research index", "I(F)=100(A−D)/(1+A+D), range −100 to +100; not truth probability."),
        ("Evidence strength", "E(F)=100(1−exp(−(A+D))); amount of weighted evidence."),
        ("Source quality", "Weighted geometric mean of six positive dimensions and bias safety."),
        ("Dependence", "Within each independence cluster only the strongest contribution per direction counts."),
        ("Stances", "+1 explicit, +0.5 implicit, 0 silence, −1.5 explicit D1 contradiction."),
        ("Disagreements", "D2–D9 remain visible but only D1 receives an automatic penalty."),
        ("Silence", "Covered-but-silent and not-covered are distinct; neither is contradiction."),
        ("Audit rule", "Every quoted passage is matched to extracted page text and carries PDF/printed pagination."),
    ]:
        method_ws.append(item)
    _configure_sheet(method_ws, {"A": 28, "B": 110}, "MethodologyTable")
    workbook.save(workbook_path)

    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in matrix:
            csv_row = list(row)
            for index in (12, 13, 14):
                csv_row[index] = round(float(csv_row[index]) * 100, 4)
            writer.writerow(csv_row)
    _write_markdown(conn, markdown_path, facts, sources, project_title)
    return ExportStats(len(facts), len(details), workbook_path, csv_path, markdown_path)
